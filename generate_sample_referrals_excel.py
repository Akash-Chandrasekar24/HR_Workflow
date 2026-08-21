import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_minimal_referral_sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Employee Referrals"

    headers = [
        "Candidate Name",
        "Candidate Email",
        "Candidate Phone",
        "Referred By",
        "Employee ID",
        "Employee Email",
        "Department",
        "Notes"
    ]

    # Minimal entries matching sample resumes in sample_resumes/
    rows = [
        [
            "Karthik Sundaram",
            "karthik.sundaram@gmail.com",
            "+91 98401 45210",
            "Alex Johnson",
            "EMP-101",
            "akashchandrasekar24@gmail.com",
            "AI Engineer",
            "Strong background in LLMs and AI system engineering."
        ],
        [
            "Sneha Venkataraman",
            "sneha.venkat@gmail.com",
            "+91 97908 12345",
            "Priya Sharma",
            "EMP-102",
            "akashchandrasekar24@gmail.com",
            "Full Stack Developer",
            "Experienced in FastAPI backend and React frontend."
        ],
        [
            "Pranav Murugan",
            "pranav.murugan@gmail.com",
            "+91 94440 67890",
            "Rahul Verma",
            "EMP-103",
            "akashchandrasekar24@gmail.com",
            "Data Analyst",
            "Skilled in SQL, Snowflake, and data warehouse modeling."
        ],
        [
            "Divya Nambiar",
            "divya.nambiar@gmail.com",
            "+91 98840 98765",
            "Dr. Vikram Seth",
            "EMP-105",
            "akashchandrasekar24@gmail.com",
            "AI Engineer",
            "Deep expertise in PyTorch and transformer models."
        ]
    ]

    # Clean GWC styling
    header_fill = PatternFill(start_color="7030B1", end_color="7030B1", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    border_thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    ws.append(headers)
    ws.row_dimensions[1].height = 26

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_thin

    for row_idx, row_data in enumerate(rows, start=2):
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 20
        bg_color = "F9FAFB" if row_idx % 2 == 0 else "FFFFFF"
        row_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Segoe UI", size=10)
            cell.fill = row_fill
            cell.border = border_thin
            cell.alignment = Alignment(
                horizontal="center" if col_idx in (3, 5, 7) else "left",
                vertical="center"
            )

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 45)

    filename = "Sample_GWC_Employee_Referrals.xlsx"
    wb.save(filename)
    print(f"Saved {filename}")

if __name__ == "__main__":
    generate_minimal_referral_sheet()

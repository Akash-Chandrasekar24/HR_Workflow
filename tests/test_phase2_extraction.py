import io
import fitz
import docx
import pytest
from app.services.document_validator import DocumentValidator
from app.services.document_extractor import DocumentExtractor
from app.services.resume_normalizer import ResumeNormalizer


def create_sample_pdf_bytes() -> bytes:
    """Generates sample PDF in-memory using PyMuPDF."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text(
        (50, 50),
        "Akash Chandrasekar\nEmail: akash@example.com\nPhone: +1-555-0199\n\n"
        "SKILLS\nPython, FastAPI, PostgreSQL, PyTorch, Docker, GCP\n\n"
        "EXPERIENCE\nAI Engineer at Tech Corp (2.5 years)\n"
        "- Developed FastAPI backend services and RAG LLM workflows.\n"
        "- Built PostgreSQL database integrations.\n\n"
        "EDUCATION\nBachelor of Science in Computer Science, 2022"
    )
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_sample_docx_bytes() -> bytes:
    """Generates sample DOCX in-memory using python-docx."""
    doc = docx.Document()
    doc.add_heading("Priya Sharma", level=1)
    doc.add_paragraph("Email: priya@example.com | Phone: +91 98765 43210")
    
    doc.add_heading("Skills", level=2)
    doc.add_paragraph("SQL, Tableau, PowerBI, Python, Pandas, Statistics")
    
    doc.add_heading("Work Experience", level=2)
    doc.add_paragraph("Data Analyst - Global Analytics Inc (3 years)")
    doc.add_paragraph("Extracted complex datasets using SQL queries and generated executive dashboards in Tableau.")
    
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_document_validator_pdf():
    pdf_bytes = create_sample_pdf_bytes()
    is_valid, fmt, sha256_hash, err = DocumentValidator.validate_and_hash(pdf_bytes, "resume.pdf")
    assert is_valid is True
    assert fmt == "PDF"
    assert len(sha256_hash) == 64
    assert err == ""


def test_document_validator_docx():
    docx_bytes = create_sample_docx_bytes()
    is_valid, fmt, sha256_hash, err = DocumentValidator.validate_and_hash(docx_bytes, "resume.docx")
    assert is_valid is True
    assert fmt == "DOCX"
    assert len(sha256_hash) == 64


def test_pdf_text_extraction():
    pdf_bytes = create_sample_pdf_bytes()
    extracted_text, method = DocumentExtractor.extract_text(pdf_bytes, "PDF")
    assert method == "NATIVE_PDF"
    assert "Akash Chandrasekar" in extracted_text
    assert "Python, FastAPI" in extracted_text


def test_docx_text_extraction():
    docx_bytes = create_sample_docx_bytes()
    extracted_text, method = DocumentExtractor.extract_text(docx_bytes, "DOCX")
    assert method == "DOCX"
    assert "Priya Sharma" in extracted_text
    assert "SQL, Tableau" in extracted_text


def test_resume_normalization_and_sections():
    raw = (
        " John Doe \n\n\n SKILLS \n Python, SQL \n\n EXPERIENCE \n Software Engineer for 2 years "
    )
    normalized = ResumeNormalizer.normalize_text(raw)
    assert "\n\n\n" not in normalized

    sections = ResumeNormalizer.extract_sections(normalized)
    assert "skills" in sections
    assert "Python, SQL" in sections["skills"]
    assert "experience" in sections

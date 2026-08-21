import io
import fitz  # PyMuPDF
import docx
import structlog
from typing import Tuple

logger = structlog.get_logger(__name__)


class DocumentExtractor:
    """
    Extracts raw text from PDF and DOCX files.
    Applies OCR fallback if native text extraction yields insufficient text (scanned PDFs).
    """

    @classmethod
    def extract_text(cls, file_bytes: bytes, file_format: str) -> Tuple[str, str]:
        """
        Extracts text from file bytes.
        Returns: (extracted_text, extraction_method)
        """
        file_format = file_format.upper()

        if file_format == "PDF":
            return cls._extract_from_pdf(file_bytes)
        elif file_format == "DOCX":
            return cls._extract_from_docx(file_bytes)
        else:
            raise ValueError(f"Unsupported file format for text extraction: {file_format}")

    @classmethod
    def _extract_from_pdf(cls, file_bytes: bytes) -> Tuple[str, str]:
        """Extracts text using PyMuPDF. Triggers OCR if text density is low."""
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        full_text = []

        for page_num in range(len(doc)):
            page = doc.load_page(page_num)
            text = page.get_text("text")
            if text.strip():
                full_text.append(text)

        combined_text = "\n".join(full_text).strip()

        # If native PDF extraction returns less than 100 characters, attempt OCR fallback
        if len(combined_text) < 100:
            logger.warning("Native PDF text yield low. Triggering OCR fallback engine.")
            ocr_text = cls._ocr_fallback_pdf(file_bytes)
            if ocr_text.strip():
                return ocr_text, "OCR"
            return combined_text, "NATIVE_PDF_LOW_QUALITY"

        return combined_text, "NATIVE_PDF"

    @classmethod
    def _extract_from_docx(cls, file_bytes: bytes) -> Tuple[str, str]:
        """Extracts text and table content from DOCX files using python-docx."""
        doc_stream = io.BytesIO(file_bytes)
        document = docx.Document(doc_stream)
        text_parts = []

        # Extract paragraph text
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text.strip())

        # Extract table text
        for table in document.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_parts.append(" | ".join(row_text))

        full_text = "\n".join(text_parts).strip()
        return full_text, "DOCX"

    @classmethod
    def _ocr_fallback_pdf(cls, file_bytes: bytes) -> str:
        """OCR fallback for scanned PDFs using pytesseract and pdf2image if available."""
        try:
            from pdf2image import convert_from_bytes
            import pytesseract

            images = convert_from_bytes(file_bytes)
            ocr_pages = []
            for image in images:
                ocr_pages.append(pytesseract.image_to_string(image))
            return "\n".join(ocr_pages)
        except Exception as e:
            logger.warning("OCR engine unavailable or failed during execution", error=str(e))
            return ""
